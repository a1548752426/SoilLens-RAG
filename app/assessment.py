from __future__ import annotations

import math
import statistics
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


METAL_ORDER = ("As", "Cd", "Cr", "Mn", "Ni", "Cu", "Pb", "Zn")
BACKGROUND_VALUES = {
    "As": 5.0,
    "Cd": 0.11,
    "Cr": 57.0,
    "Mn": 465.0,
    "Ni": 26.0,
    "Cu": 16.0,
    "Pb": 21.0,
    "Zn": 68.0,
}
SOURCE_TITLE = "中国城市土壤化学元素的背景值与基准值"
SOURCE_TABLE = "表 12：杭州市土壤地球化学背景值与基准值"
SOURCE_PAGE = 16
UNIT = "mg/kg"


def classify_cf(value: float) -> str:
    if value < 1:
        return "低污染"
    if value < 3:
        return "中等污染"
    if value < 6:
        return "较高污染"
    return "极高污染"


def classify_igeo(value: float | None) -> str:
    if value is None:
        return "无法计算"
    if value < 0:
        return "无污染"
    if value < 1:
        return "无污染至中度污染"
    if value < 2:
        return "中度污染"
    if value < 3:
        return "中度至强污染"
    if value < 4:
        return "强污染"
    if value < 5:
        return "强至极强污染"
    return "极强污染"


def classify_pli(value: float) -> str:
    if math.isclose(value, 1.0, rel_tol=1e-9, abs_tol=1e-9):
        return "背景水平"
    if value < 1:
        return "整体未污染"
    return "整体污染"


def calculate_assessment(
    concentrations: dict[str, float],
    *,
    sample_id: str | None = None,
    sample_name: str | None = None,
) -> dict[str, Any]:
    available: dict[str, float] = {}
    for metal in METAL_ORDER:
        if metal not in concentrations:
            continue
        value = float(concentrations[metal])
        if not math.isfinite(value) or value < 0:
            raise ValueError(f"{metal} 浓度必须是大于或等于 0 的有限数值。")
        available[metal] = value

    if not available:
        raise ValueError("至少需要输入一种重金属浓度。")

    rows: list[dict[str, Any]] = []
    factors: list[float] = []
    for metal in METAL_ORDER:
        if metal not in available:
            continue
        concentration = available[metal]
        background = BACKGROUND_VALUES[metal]
        cf = concentration / background
        igeo = math.log2(concentration / (1.5 * background)) if concentration > 0 else None
        factors.append(cf)
        rows.append(
            {
                "metal": metal,
                "concentration": concentration,
                "background": background,
                "cf": cf,
                "cf_grade": classify_cf(cf),
                "igeo": igeo,
                "igeo_grade": classify_igeo(igeo),
                "exceeded": concentration > background,
            }
        )

    if any(factor == 0 for factor in factors):
        pli = 0.0
    else:
        pli = math.exp(sum(math.log(factor) for factor in factors) / len(factors))

    max_cf = max(rows, key=lambda row: row["cf"])
    calculable_igeo = [row for row in rows if row["igeo"] is not None]
    max_igeo = max(calculable_igeo, key=lambda row: row["igeo"]) if calculable_igeo else None
    exceeded_count = sum(bool(row["exceeded"]) for row in rows)
    metals = [row["metal"] for row in rows]

    warnings = [
        "本结果表示相对杭州市城市土壤地球化学背景值的污染程度，不等同于农用地或建设用地风险筛选结论。"
    ]
    missing = [metal for metal in METAL_ORDER if metal not in available]
    if missing:
        missing_prefix = "当前自定义样点未填写" if sample_id is None else "当前数据未提供"
        warnings.insert(
            0,
            f"{missing_prefix} {'、'.join(missing)}，PLI 按现有 {len(metals)} 种金属计算。",
        )

    return {
        "sample": {
            "sample_id": sample_id,
            "name": sample_name,
            "concentrations": available,
        },
        "summary": {
            "pli": pli,
            "pli_grade": classify_pli(pli),
            "metal_count": len(rows),
            "exceeded_count": exceeded_count,
            "max_cf_metal": max_cf["metal"],
            "max_cf": max_cf["cf"],
            "max_igeo_metal": max_igeo["metal"] if max_igeo else None,
            "max_igeo": max_igeo["igeo"] if max_igeo else None,
        },
        "results": rows,
        "metals": metals,
        "warnings": warnings,
        "method": {
            "cf": "CF = C / B",
            "igeo": "Igeo = log2[C / (1.5 × B)]",
            "pli": f"PLI = (CF₁ × CF₂ × … × CFₙ)^(1/n)，本次 n = {len(rows)}",
        },
    }


def calculate_batch_assessment(
    samples: list[dict[str, Any]],
    *,
    include_details: bool = False,
) -> dict[str, Any]:
    if not samples:
        raise ValueError("没有可用于批量评价的样点数据。")

    sample_rows: list[dict[str, Any]] = []
    details: list[dict[str, Any]] = []
    metal_accumulators = {
        metal: {"rows": [], "exceeded_count": 0}
        for metal in METAL_ORDER
    }
    pli_grade_distribution = {
        "整体未污染": 0,
        "背景水平": 0,
        "整体污染": 0,
    }

    for sample in samples:
        sample_id = str(sample.get("sample_id", ""))
        name = str(sample.get("name") or f"样点 {sample_id}")
        concentrations = dict(sample.get("concentrations") or {})
        assessment = calculate_assessment(
            concentrations,
            sample_id=sample_id,
            sample_name=name,
        )
        summary = assessment["summary"]
        pli_grade_distribution[summary["pli_grade"]] += 1
        sample_rows.append(
            {
                "sample_id": sample_id,
                "name": name,
                "pli": summary["pli"],
                "pli_grade": summary["pli_grade"],
                "metal_count": summary["metal_count"],
                "exceeded_count": summary["exceeded_count"],
                "max_cf_metal": summary["max_cf_metal"],
                "max_cf": summary["max_cf"],
                "max_igeo_metal": summary["max_igeo_metal"],
                "max_igeo": summary["max_igeo"],
            }
        )
        for row in assessment["results"]:
            metal = row["metal"]
            metal_accumulators[metal]["rows"].append(row)
            if row["exceeded"]:
                metal_accumulators[metal]["exceeded_count"] += 1
            if include_details:
                details.append(
                    {
                        "sample_id": sample_id,
                        "name": name,
                        **row,
                    }
                )

    sample_rows.sort(key=lambda row: row["pli"], reverse=True)
    metal_summary: list[dict[str, Any]] = []
    for metal in METAL_ORDER:
        accumulator = metal_accumulators[metal]
        rows = accumulator["rows"]
        if not rows:
            continue
        count = len(rows)
        exceeded_count = accumulator["exceeded_count"]
        igeo_values = [row["igeo"] for row in rows if row["igeo"] is not None]
        metal_summary.append(
            {
                "metal": metal,
                "background": BACKGROUND_VALUES[metal],
                "sample_count": count,
                "exceeded_count": exceeded_count,
                "exceeded_rate": exceeded_count / count,
                "mean_cf": statistics.fmean(row["cf"] for row in rows),
                "max_cf": max(row["cf"] for row in rows),
                "max_igeo": max(igeo_values) if igeo_values else None,
                "igeo_polluted_count": sum(
                    row["igeo"] is not None and row["igeo"] >= 0
                    for row in rows
                ),
            }
        )

    pli_values = [row["pli"] for row in sample_rows]
    polluted_count = sum(row["pli"] > 1 for row in sample_rows)
    max_pli_sample = sample_rows[0]
    highest_exceedance = max(metal_summary, key=lambda row: row["exceeded_rate"])
    summary = {
        "sample_count": len(sample_rows),
        "metal_count": len(metal_summary),
        "polluted_count": polluted_count,
        "polluted_rate": polluted_count / len(sample_rows),
        "mean_pli": statistics.fmean(pli_values),
        "median_pli": statistics.median(pli_values),
        "max_pli": max_pli_sample["pli"],
        "max_pli_sample_id": max_pli_sample["sample_id"],
        "max_pli_sample_name": max_pli_sample["name"],
        "highest_exceedance_metal": highest_exceedance["metal"],
        "highest_exceedance_rate": highest_exceedance["exceeded_rate"],
    }

    result = {
        "summary": summary,
        "pli_grade_distribution": pli_grade_distribution,
        "metal_summary": metal_summary,
        "samples": sample_rows,
        "top_samples": sample_rows[:10],
        "note": (
            "批量结果仅在本机生成，表示相对杭州市城市土壤地球化学背景值的污染程度，"
            "不等同于农用地或建设用地风险筛选结论。"
        ),
    }
    if include_details:
        result["details"] = details
    return result


@dataclass(frozen=True)
class SoilSample:
    sample_id: str
    name: str
    concentrations: dict[str, float]

    def public(self) -> dict[str, Any]:
        return {
            "sample_id": self.sample_id,
            "name": self.name,
            "concentrations": self.concentrations,
        }


class AssessmentDataset:
    def __init__(
        self,
        workbook_path: Path,
        source_pdf_path: Path,
        *,
        synthetic: bool = False,
    ):
        self.workbook_path = workbook_path
        self.source_pdf_path = source_pdf_path
        self.synthetic = synthetic
        self._samples: list[SoilSample] | None = None
        self._samples_by_id: dict[str, SoilSample] = {}

    def _load(self) -> None:
        if not self.workbook_path.exists():
            self._samples = []
            self._samples_by_id = {}
            return

        workbook = load_workbook(self.workbook_path, read_only=True, data_only=True)
        try:
            sheet = workbook.active
            rows = sheet.iter_rows(values_only=True)
            raw_headers = next(rows, None)
            if not raw_headers:
                raise ValueError("样点数据表为空。")
            headers = [str(value).strip() if value is not None else "" for value in raw_headers]
            required = {"Sample_ID", "name"}
            missing_required = required.difference(headers)
            if missing_required:
                raise ValueError(f"样点数据缺少字段：{', '.join(sorted(missing_required))}")

            metal_headers = [metal for metal in METAL_ORDER if metal in headers]
            if not metal_headers:
                raise ValueError("样点数据中未找到可识别的重金属字段。")

            positions = {header: index for index, header in enumerate(headers)}
            samples: list[SoilSample] = []
            seen: set[str] = set()
            for row_number, values in enumerate(rows, start=2):
                raw_id = values[positions["Sample_ID"]]
                if raw_id is None:
                    continue
                sample_id = str(raw_id).strip()
                if sample_id.endswith(".0"):
                    sample_id = sample_id[:-2]
                if sample_id in seen:
                    raise ValueError(f"Sample_ID 重复：{sample_id}")
                seen.add(sample_id)

                raw_name = values[positions["name"]]
                name = str(raw_name).strip() if raw_name is not None else f"样点 {sample_id}"
                concentrations: dict[str, float] = {}
                for metal in metal_headers:
                    raw_value = values[positions[metal]]
                    if raw_value is None:
                        continue
                    try:
                        value = float(raw_value)
                    except (TypeError, ValueError) as exc:
                        raise ValueError(f"第 {row_number} 行 {metal} 不是有效数值。") from exc
                    if not math.isfinite(value) or value < 0:
                        raise ValueError(f"第 {row_number} 行 {metal} 必须大于或等于 0。")
                    concentrations[metal] = value

                samples.append(SoilSample(sample_id, name, concentrations))
        finally:
            workbook.close()

        self._samples = samples
        self._samples_by_id = {sample.sample_id: sample for sample in samples}

    def list_samples(self) -> list[dict[str, Any]]:
        if self._samples is None:
            self._load()
        return [sample.public() for sample in self._samples or []]

    def get_sample(self, sample_id: str) -> SoilSample | None:
        if self._samples is None:
            self._load()
        return self._samples_by_id.get(str(sample_id))

    def config(self) -> dict[str, Any]:
        samples = self.list_samples()
        dataset_available = bool(samples)
        available_metals = (
            [
                metal
                for metal in METAL_ORDER
                if any(metal in sample["concentrations"] for sample in samples)
            ]
            if dataset_available
            else list(METAL_ORDER)
        )
        missing_metals = [metal for metal in METAL_ORDER if metal not in available_metals]
        data_label = "合成演示数据" if self.synthetic else "私有数据"
        if not dataset_available:
            note = f"未加载{data_label}文件。可以手动输入浓度，输入内容仅用于本次计算。"
        elif missing_metals:
            note = (
                f"{data_label}的 Excel 未提供 {'、'.join(missing_metals)}，样点 PLI 将按 "
                f"{'、'.join(available_metals)} 共 {len(available_metals)} 种金属计算。"
            )
        else:
            note = f"{data_label}的样点 PLI 按 {len(available_metals)} 种金属计算。"
        if self.synthetic:
            note += " 所有浓度均为人工合成，仅用于功能演示，不可用于科研或监管结论。"
        return {
            "dataset_available": dataset_available,
            "synthetic": self.synthetic,
            "sample_count": len(samples),
            "available_metals": available_metals,
            "missing_metals": missing_metals,
            "background_values": BACKGROUND_VALUES,
            "unit": UNIT,
            "source": {
                "title": SOURCE_TITLE,
                "table": SOURCE_TABLE,
                "page": SOURCE_PAGE,
                "url": f"/api/assessment/source#page={SOURCE_PAGE}",
            },
            "formulas": {
                "cf": "CF = C / B",
                "igeo": "Igeo = log2[C / (1.5 × B)]",
                "pli": "PLI = (CF₁ × CF₂ × … × CFₙ)^(1/n)",
            },
            "note": note,
        }
