from __future__ import annotations

from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .assessment import (
    BACKGROUND_VALUES,
    SOURCE_PAGE,
    SOURCE_TABLE,
    SOURCE_TITLE,
    UNIT,
)


GREEN = "28765B"
DARK_GREEN = "173A2D"
PALE_GREEN = "EAF3E5"
PALE_GOLD = "FFF2D8"
PALE_RED = "F9DDDA"
LIGHT_LINE = "DDE2D9"
WHITE = "FFFFFF"
MUTED = "66756C"


def _style_title(sheet, title: str, end_column: int) -> None:
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=end_column)
    cell = sheet.cell(1, 1, title)
    cell.fill = PatternFill("solid", fgColor=DARK_GREEN)
    cell.font = Font(color=WHITE, bold=True, size=16)
    cell.alignment = Alignment(vertical="center")
    sheet.row_dimensions[1].height = 30
    sheet.sheet_view.showGridLines = False


def _style_header(sheet, row_number: int, column_count: int) -> None:
    for cell in sheet[row_number][:column_count]:
        cell.fill = PatternFill("solid", fgColor=GREEN)
        cell.font = Font(color=WHITE, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(bottom=Side(style="thin", color=LIGHT_LINE))


def _fit_columns(sheet, widths: dict[int, float]) -> None:
    for column, width in widths.items():
        sheet.column_dimensions[get_column_letter(column)].width = width


def build_batch_workbook(batch: dict[str, Any]) -> bytes:
    workbook = Workbook()
    overview = workbook.active
    overview.title = "总览"

    _style_title(overview, "SoilLens 土壤重金属批量污染评价", 8)
    overview["A3"] = "指标"
    overview["B3"] = "结果"
    _style_header(overview, 3, 2)
    summary = batch["summary"]
    overview_rows = [
        ("样点数量", summary["sample_count"]),
        ("参与元素数量", summary["metal_count"]),
        ("PLI > 1 样点数", summary["polluted_count"]),
        ("PLI > 1 比例", summary["polluted_rate"]),
        ("平均 PLI", summary["mean_pli"]),
        ("中位数 PLI", summary["median_pli"]),
        ("最高 PLI 样点", f"{summary['max_pli_sample_id']} · {summary['max_pli_sample_name']}"),
        ("最高 PLI", summary["max_pli"]),
        ("超背景率最高元素", summary["highest_exceedance_metal"]),
        ("该元素超背景率", summary["highest_exceedance_rate"]),
    ]
    for row_index, values in enumerate(overview_rows, start=4):
        overview.cell(row_index, 1, values[0])
        overview.cell(row_index, 2, values[1])
    overview["B7"].number_format = "0.0%"
    overview["B8"].number_format = "0.000"
    overview["B9"].number_format = "0.000"
    overview["B11"].number_format = "0.000"
    overview["B13"].number_format = "0.0%"

    metal_header_row = 16
    metal_headers = [
        "元素",
        f"背景值 ({UNIT})",
        "样点数",
        "超背景数",
        "超背景率",
        "平均 CF",
        "最大 CF",
        "最大 Igeo",
    ]
    for column, value in enumerate(metal_headers, start=1):
        overview.cell(metal_header_row, column, value)
    _style_header(overview, metal_header_row, len(metal_headers))
    for row_index, row in enumerate(batch["metal_summary"], start=metal_header_row + 1):
        values = [
            row["metal"],
            row["background"],
            row["sample_count"],
            row["exceeded_count"],
            row["exceeded_rate"],
            row["mean_cf"],
            row["max_cf"],
            row["max_igeo"],
        ]
        for column, value in enumerate(values, start=1):
            overview.cell(row_index, column, value)
        overview.cell(row_index, 5).number_format = "0.0%"
        for column in (2, 6, 7, 8):
            overview.cell(row_index, column).number_format = "0.000"

    note_row = metal_header_row + len(batch["metal_summary"]) + 3
    overview.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=8)
    overview.cell(note_row, 1, batch["note"])
    overview.cell(note_row, 1).fill = PatternFill("solid", fgColor=PALE_GOLD)
    overview.cell(note_row, 1).font = Font(color=MUTED)
    overview.cell(note_row, 1).alignment = Alignment(wrap_text=True, vertical="center")
    overview.row_dimensions[note_row].height = 36
    _fit_columns(overview, {1: 22, 2: 24, 3: 13, 4: 13, 5: 14, 6: 14, 7: 14, 8: 15})
    overview.freeze_panes = "A4"

    samples_sheet = workbook.create_sheet("样点汇总")
    _style_title(samples_sheet, "样点综合污染评价汇总", 10)
    sample_headers = [
        "Sample_ID",
        "样点名称",
        "PLI",
        "PLI 分级",
        "超背景元素数",
        "参与元素数",
        "最高 CF 元素",
        "最高 CF",
        "最高 Igeo 元素",
        "最高 Igeo",
    ]
    for column, value in enumerate(sample_headers, start=1):
        samples_sheet.cell(3, column, value)
    _style_header(samples_sheet, 3, len(sample_headers))
    for row_index, row in enumerate(batch["samples"], start=4):
        values = [
            row["sample_id"],
            row["name"],
            row["pli"],
            row["pli_grade"],
            row["exceeded_count"],
            row["metal_count"],
            row["max_cf_metal"],
            row["max_cf"],
            row["max_igeo_metal"],
            row["max_igeo"],
        ]
        for column, value in enumerate(values, start=1):
            samples_sheet.cell(row_index, column, value)
        for column in (3, 8, 10):
            samples_sheet.cell(row_index, column).number_format = "0.000"
    final_sample_row = 3 + len(batch["samples"])
    samples_sheet.auto_filter.ref = f"A3:J{final_sample_row}"
    samples_sheet.freeze_panes = "A4"
    samples_sheet.conditional_formatting.add(
        f"C4:C{final_sample_row}",
        CellIsRule(
            operator="greaterThan",
            formula=["1"],
            fill=PatternFill("solid", fgColor=PALE_RED),
        ),
    )
    _fit_columns(
        samples_sheet,
        {1: 14, 2: 28, 3: 12, 4: 17, 5: 16, 6: 14, 7: 15, 8: 12, 9: 17, 10: 13},
    )

    details_sheet = workbook.create_sheet("逐元素结果")
    _style_title(details_sheet, "样点逐元素污染评价结果", 10)
    detail_headers = [
        "Sample_ID",
        "样点名称",
        "元素",
        f"浓度 ({UNIT})",
        f"背景值 ({UNIT})",
        "CF",
        "CF 分级",
        "Igeo",
        "Igeo 分级",
        "超过背景值",
    ]
    for column, value in enumerate(detail_headers, start=1):
        details_sheet.cell(3, column, value)
    _style_header(details_sheet, 3, len(detail_headers))
    for row_index, row in enumerate(batch.get("details", []), start=4):
        values = [
            row["sample_id"],
            row["name"],
            row["metal"],
            row["concentration"],
            row["background"],
            row["cf"],
            row["cf_grade"],
            row["igeo"],
            row["igeo_grade"],
            "是" if row["exceeded"] else "否",
        ]
        for column, value in enumerate(values, start=1):
            details_sheet.cell(row_index, column, value)
        for column in (4, 5, 6, 8):
            details_sheet.cell(row_index, column).number_format = "0.000"
    final_detail_row = 3 + len(batch.get("details", []))
    details_sheet.auto_filter.ref = f"A3:J{final_detail_row}"
    details_sheet.freeze_panes = "A4"
    _fit_columns(
        details_sheet,
        {1: 14, 2: 28, 3: 10, 4: 16, 5: 17, 6: 12, 7: 20, 8: 12, 9: 22, 10: 15},
    )

    method_sheet = workbook.create_sheet("计算说明")
    _style_title(method_sheet, "计算方法与数据来源", 2)
    method_rows = [
        ("项目", "说明"),
        ("CF", "CF = C / B"),
        ("Igeo", "Igeo = log2[C / (1.5 × B)]"),
        ("PLI", "PLI = (CF₁ × CF₂ × … × CFₙ)^(1/n)"),
        ("C", "样点重金属浓度"),
        ("B", "杭州市城市土壤地球化学背景值"),
        ("来源", SOURCE_TITLE),
        ("表格", SOURCE_TABLE),
        ("PDF 页码", SOURCE_PAGE),
        ("使用边界", batch["note"]),
    ]
    for row_index, values in enumerate(method_rows, start=3):
        method_sheet.cell(row_index, 1, values[0])
        method_sheet.cell(row_index, 2, values[1])
    method_sheet["B12"] = (
        "本表仅在本机生成，按杭州市城市土壤地球化学背景值评价。\n"
        "不等同于农用地或建设用地风险筛选结论。"
    )
    _style_header(method_sheet, 3, 2)
    for row_index in range(4, 13):
        method_sheet.cell(row_index, 2).alignment = Alignment(
            horizontal="left",
            vertical="center",
        )
    background_row = 15
    method_sheet.cell(background_row, 1, "元素")
    method_sheet.cell(background_row, 2, f"背景值 ({UNIT})")
    _style_header(method_sheet, background_row, 2)
    for row_index, (metal, value) in enumerate(BACKGROUND_VALUES.items(), start=background_row + 1):
        method_sheet.cell(row_index, 1, metal)
        method_sheet.cell(row_index, 2, value)
        method_sheet.cell(row_index, 2).number_format = "0.000"
    method_sheet.column_dimensions["A"].width = 20
    method_sheet.column_dimensions["B"].width = 64
    method_sheet["B12"].alignment = Alignment(wrap_text=True)
    method_sheet.row_dimensions[12].height = 42

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
