from pathlib import Path
from io import BytesIO
import math
import tempfile
import unittest

from openpyxl import Workbook, load_workbook

from app.assessment import (
    AssessmentDataset,
    BACKGROUND_VALUES,
    calculate_assessment,
    calculate_batch_assessment,
)
from app.exporter import build_batch_workbook


ROOT = Path(__file__).resolve().parents[1]


class AssessmentTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temp_directory = tempfile.TemporaryDirectory()
        workbook_path = Path(cls.temp_directory.name) / "synthetic_samples.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["Sample_ID", "name", "As", "Cd", "Cr", "Mn", "Ni", "Cu", "Zn"])
        sheet.append(["demo-1", "合成样点 A", 10, 0.22, 57, 465, 26, 16, 68])
        sheet.append(["demo-2", "合成样点 B", 5, 0.11, 28.5, 232.5, 13, 8, 34])
        workbook.save(workbook_path)
        workbook.close()
        cls.dataset = AssessmentDataset(
            workbook_path,
            ROOT / "background_values" / "中国城市土壤化学元素的背景值与基准值_成杭新.pdf",
        )

    @classmethod
    def tearDownClass(cls):
        cls.temp_directory.cleanup()

    def test_synthetic_workbook_contains_two_samples_and_seven_metals(self):
        config = self.dataset.config()
        self.assertEqual(config["sample_count"], 2)
        self.assertEqual(
            config["available_metals"],
            ["As", "Cd", "Cr", "Mn", "Ni", "Cu", "Zn"],
        )
        self.assertEqual(config["missing_metals"], ["Pb"])

    def test_background_values_match_table_12(self):
        self.assertEqual(
            BACKGROUND_VALUES,
            {
                "As": 5.0,
                "Cd": 0.11,
                "Cr": 57.0,
                "Mn": 465.0,
                "Ni": 26.0,
                "Cu": 16.0,
                "Pb": 21.0,
                "Zn": 68.0,
            },
        )

    def test_first_sample_calculation_is_reproducible(self):
        sample = self.dataset.get_sample("demo-1")
        self.assertIsNotNone(sample)
        result = calculate_assessment(
            sample.concentrations,
            sample_id=sample.sample_id,
            sample_name=sample.name,
        )

        expected_cf = sample.concentrations["As"] / BACKGROUND_VALUES["As"]
        expected_igeo = math.log2(
            sample.concentrations["As"] / (1.5 * BACKGROUND_VALUES["As"])
        )
        factors = [
            sample.concentrations[metal] / BACKGROUND_VALUES[metal]
            for metal in result["metals"]
        ]
        expected_pli = math.prod(factors) ** (1 / len(factors))

        as_row = next(row for row in result["results"] if row["metal"] == "As")
        self.assertAlmostEqual(as_row["cf"], expected_cf, places=12)
        self.assertAlmostEqual(as_row["igeo"], expected_igeo, places=12)
        self.assertAlmostEqual(result["summary"]["pli"], expected_pli, places=12)
        self.assertEqual(result["summary"]["metal_count"], 7)
        self.assertIn("Pb", result["warnings"][0])

    def test_zero_concentration_does_not_crash(self):
        result = calculate_assessment({"As": 0, "Cd": 0.11})
        self.assertEqual(result["summary"]["pli"], 0)
        as_row = next(row for row in result["results"] if row["metal"] == "As")
        self.assertIsNone(as_row["igeo"])

    def test_custom_sample_can_use_one_metal(self):
        result = calculate_assessment({"Pb": 42}, sample_name="自定义样点")
        self.assertEqual(result["summary"]["metal_count"], 1)
        self.assertEqual(result["metals"], ["Pb"])
        self.assertAlmostEqual(result["summary"]["pli"], 2.0)
        self.assertEqual(result["results"][0]["cf"], 2.0)
        self.assertIn("现有 1 种金属", result["warnings"][0])

    def test_batch_assessment_summarizes_and_sorts_samples(self):
        batch = calculate_batch_assessment(
            self.dataset.list_samples(),
            include_details=True,
        )
        self.assertEqual(batch["summary"]["sample_count"], 2)
        self.assertEqual(batch["summary"]["metal_count"], 7)
        self.assertEqual(len(batch["metal_summary"]), 7)
        self.assertEqual(len(batch["details"]), 14)
        self.assertGreaterEqual(batch["samples"][0]["pli"], batch["samples"][1]["pli"])

    def test_batch_export_contains_auditable_sheets(self):
        batch = calculate_batch_assessment(
            self.dataset.list_samples(),
            include_details=True,
        )
        content = build_batch_workbook(batch)
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=False)
        try:
            self.assertEqual(
                workbook.sheetnames,
                ["总览", "样点汇总", "逐元素结果", "计算说明"],
            )
            self.assertEqual(workbook["样点汇总"].max_row, 5)
            self.assertEqual(workbook["逐元素结果"].max_row, 17)
            self.assertEqual(workbook["计算说明"]["B4"].value, "CF = C / B")
            self.assertIn("\n不等同", workbook["计算说明"]["B12"].value)
        finally:
            workbook.close()

    def test_missing_private_workbook_uses_manual_mode(self):
        dataset = AssessmentDataset(
            Path(self.temp_directory.name) / "not_present.xlsx",
            ROOT / "background_values" / "中国城市土壤化学元素的背景值与基准值_成杭新.pdf",
        )
        config = dataset.config()
        self.assertFalse(config["dataset_available"])
        self.assertEqual(config["sample_count"], 0)
        self.assertEqual(dataset.list_samples(), [])


if __name__ == "__main__":
    unittest.main()
